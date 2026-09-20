"""Real PPTX extraction through Switch Bay ce_ingest → CE local_ingest."""

from __future__ import annotations

import os
import sqlite3
import subprocess
import sys
from pathlib import Path

import pytest
from pptx import Presentation
from pptx.util import Inches

from switchbay import ce_tools, cebridge, tools, watchfolders

PPTX_TITLE = "SBX-PPTX-Title-9f3c2a17"
PPTX_BULLET_A = "SBX-PPTX-Bullet-ALPHA-9f3c2a17"
PPTX_BULLET_B = "SBX-PPTX-Bullet-BETA-9f3c2a17"
PPTX_CELL = "ZX9-TABLE-CELL-9f3c2a17"
PPTX_METRIC = "UniqueKey-9f3c2a17"


def _ce_ingest_script() -> Path:
    return cebridge.ce_root() / "scripts" / "local_ingest.py"


def _require_ce() -> Path:
    script = _ce_ingest_script()
    if not script.is_file():
        pytest.skip("optional curiosity-engine installation absent")
    return script


def _write_pptx(path: Path) -> None:
    prs = Presentation()
    layout = prs.slide_layouts[1]  # title + body
    slide = prs.slides.add_slide(layout)
    slide.shapes.title.text = PPTX_TITLE
    body = slide.placeholders[1].text_frame
    body.text = PPTX_BULLET_A
    p = body.add_paragraph()
    p.text = PPTX_BULLET_B
    p.level = 0

    blank = prs.slide_layouts[6] if len(prs.slide_layouts) > 6 else prs.slide_layouts[5]
    table_slide = prs.slides.add_slide(blank)
    rows, cols = 3, 2
    table = table_slide.shapes.add_table(
        rows, cols, Inches(0.5), Inches(1.0), Inches(8.0), Inches(2.0),
    ).table
    table.cell(0, 0).text = "Metric"
    table.cell(0, 1).text = "Value"
    table.cell(1, 0).text = PPTX_METRIC
    table.cell(1, 1).text = PPTX_CELL
    table.cell(2, 0).text = "Count"
    table.cell(2, 1).text = "42"
    path.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(path))


def _workspace(tmp_path: Path) -> Path:
    ws = tmp_path / "ws"
    (ws / "vault" / "raw").mkdir(parents=True)
    (ws / "wiki").mkdir()
    return ws


def _extracted_texts(ws: Path) -> list[str]:
    return [
        p.read_text(encoding="utf-8", errors="replace")
        for p in sorted((ws / "vault").rglob("*.extracted.md"))
    ]


def test_host_python_has_pptx_and_workspace_venv_can_lack_it(tmp_path: Path):
    assert cebridge.host_has_module("pptx")
    info = cebridge.host_extractor_info()
    assert info["executable"] == sys.executable
    assert info["modules"]["pptx"]
    ws = _workspace(tmp_path)
    subprocess.run(
        [sys.executable, "-m", "venv", "--without-pip", str(ws / ".venv")],
        check=True,
    )
    py = ws / ".venv" / ("Scripts/python.exe" if os.name == "nt" else "bin/python")
    ingest_py = cebridge.script_python_for(ws, "local_ingest.py", ext=".pptx")
    graph_py = cebridge.script_python_for(ws, "graph.py")
    scan_py = cebridge.script_python_for(ws, "scan.py")
    xlsx_py = cebridge.python_for_ingest(ws, ".xlsx")
    assert ingest_py == [sys.executable]
    assert graph_py == [str(py)]
    assert scan_py == [str(py)]
    assert xlsx_py == [str(py)]
    assert ingest_py != graph_py


def test_real_ce_pptx_ingest_extracts_title_bullets_table(tmp_path: Path):
    _require_ce()
    ws = _workspace(tmp_path)
    src = ws / "vault" / "raw" / "sbx-pptx-9f3c2a17.pptx"
    _write_pptx(src)
    interp = cebridge.script_python_for(ws, "local_ingest.py", ext=".pptx")
    assert interp == [sys.executable]
    out = tools.REGISTRY["ce_ingest"].handler(
        ws, {"path": "vault/raw/sbx-pptx-9f3c2a17.pptx"},
    )
    assert ce_tools.ingest_is_success(out), out
    texts = _extracted_texts(ws)
    assert texts, f"no extracted.md written: {out}"
    body = "\n".join(texts)
    assert PPTX_TITLE in body
    assert PPTX_BULLET_A in body
    assert PPTX_BULLET_B in body
    assert PPTX_CELL in body
    assert PPTX_METRIC in body
    assert "extraction unavailable" not in body.lower()
    assert "python-pptx` not" not in body
    methods = []
    rows = out.get("results") if isinstance(out.get("results"), list) else [out]
    for row in rows:
        if isinstance(row, dict) and row.get("extraction_method"):
            methods.append(row["extraction_method"])
    assert methods, out
    assert all(m == "python-pptx" for m in methods), methods
    used = out.get("interpreter")
    if isinstance(used, list) and used and isinstance(used[0], list):
        used = used[0]
    assert used == [sys.executable], used


def test_real_ce_txt_ingest_still_works(tmp_path: Path):
    _require_ce()
    ws = _workspace(tmp_path)
    note = ws / "vault" / "raw" / "plain-9f3c2a17.txt"
    note.write_text("SBX-TXT-BODY-9f3c2a17\nsecond line\n", encoding="utf-8")
    out = tools.REGISTRY["ce_ingest"].handler(
        ws, {"path": "vault/raw/plain-9f3c2a17.txt"},
    )
    assert ce_tools.ingest_is_success(out), out
    body = "\n".join(_extracted_texts(ws))
    assert "SBX-TXT-BODY-9f3c2a17" in body


def test_corrupt_pptx_is_retryable_failure_not_placeholder_success(tmp_path: Path):
    _require_ce()
    ws = _workspace(tmp_path)
    bad = ws / "vault" / "raw" / "corrupt-9f3c2a17.pptx"
    bad.write_bytes(b"PK\x03\x04 this is not a real presentation")
    out = tools.REGISTRY["ce_ingest"].handler(
        ws, {"path": "vault/raw/corrupt-9f3c2a17.pptx"},
    )
    assert not ce_tools.ingest_is_success(out), out
    assert out.get("ok") is False
    assert out.get("retryable") is True
    extracts = list((ws / "vault").rglob("*.extracted.md"))
    assert extracts == [], f"placeholder extract accepted: {extracts}"
    assert bad.is_file()
    db = ws / "vault" / "vault.db"
    if db.is_file():
        with sqlite3.connect(db) as conn:
            paths = [r[0] for r in conn.execute("SELECT path FROM sources")]
        assert not any("corrupt-9f3c2a17" in p for p in paths), paths


def test_watch_handoff_runs_real_pptx_extract(tmp_path: Path, monkeypatch: pytest.MonkeyPatch):
    _require_ce()
    monkeypatch.setattr(
        "switchbay.workspaces.is_within_home", lambda _p: True,
    )
    monkeypatch.setattr(watchfolders, "SETTLE_SECONDS", 0.0)
    ws = _workspace(tmp_path)
    folder = tmp_path / "watch"
    folder.mkdir()
    rec = watchfolders.add_folder(ws, str(folder))
    assert isinstance(rec, dict)
    src = folder / "watched-9f3c2a17.pptx"
    _write_pptx(src)
    picked, _backlog = watchfolders.scan_candidates(ws)
    assert [c.logical for c in picked] == [str(src.resolve())]
    seen_before = watchfolders._load_seen(ws)
    assert str(src.resolve()) not in seen_before
    result = watchfolders.handoff(ws, picked[0])
    assert result.status == "success", result
    assert result.extracted
    seen_after = watchfolders._load_seen(ws)
    assert str(src.resolve()) in seen_after
    body = "\n".join(_extracted_texts(ws))
    assert PPTX_TITLE in body
    assert PPTX_CELL in body
    assert "extraction unavailable" not in body.lower()
    orig = str(src.resolve())
    assert f"source_path: {orig}" in body
    assert f"extracted_from: {orig}" in body
    # Second beat does not re-hand-off.
    picked2, _ = watchfolders.scan_candidates(ws)
    assert picked2 == []


# Optional developer venv with real openpyxl/pypdf. Clean CI does not
# have this path; capability-selection builds a stub venv instead.
FORMAT_VENV = Path("/tmp/switchbay-format-regression/.venv")


def _venv_python(venv: Path) -> Path | None:
    for rel in (
        Path("bin") / "python",
        Path("bin") / "python3",
        Path("Scripts") / "python.exe",
    ):
        p = venv / rel
        if p.is_file():
            return p
    return None


def _purelib_of(py: Path) -> Path:
    probe = subprocess.run(
        [str(py), "-c", "import sysconfig; print(sysconfig.get_path('purelib'))"],
        check=True,
        capture_output=True,
        text=True,
    )
    lines = [ln.strip() for ln in probe.stdout.splitlines() if ln.strip()]
    assert lines, probe.stdout
    path = Path(lines[-1])
    assert path.is_dir(), (path, probe.stdout, probe.stderr)
    return path


def _write_xlsx(path: Path, py: Path) -> None:
    subprocess.run(
        [
            str(py), "-c",
            "from openpyxl import Workbook; import sys; w=Workbook(); "
            "s=w.active; s.append(['XLSX-ROOT-MARKER', 'Value']); "
            "s.append(['table survived', 731]); w.save(sys.argv[1])",
            str(path),
        ],
        check=True,
    )


def _write_text_pdf(path: Path) -> None:
    content = (
        b"BT /F1 12 Tf 50 750 Td ("
        + b"PDF-ROOT-MARKER "
        + b"This synthetic document verifies ordinary readable text extraction "
          b"and preservation of existing interpreter capabilities. " * 12
        + b") Tj ET"
    )
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
        b"/Resources << /Font << /F1 4 0 R >> >> /Contents 5 0 R >>",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
        b"<< /Length " + str(len(content)).encode() + b" >>\nstream\n"
        + content + b"\nendstream",
    ]
    data = b"%PDF-1.4\n"
    offsets = [0]
    for n, obj in enumerate(objects, 1):
        offsets.append(len(data))
        data += f"{n} 0 obj\n".encode() + obj + b"\nendobj\n"
    xref = len(data)
    data += (
        b"xref\n0 6\n0000000000 65535 f \n"
        + b"".join(f"{o:010d} 00000 n \n".encode() for o in offsets[1:])
        + f"trailer\n<< /Size 6 /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode()
    )
    path.write_bytes(data)


def test_xlsx_pdf_prefer_workspace_venv_that_has_extractors(tmp_path: Path):
    """Capability selection: workspace wins for PDF/XLSX when it can import.

    Builds a bare temp venv, drops importable pypdf/openpyxl stubs in its
    own purelib, leaves pptx absent. Probes the real interpreter — no
    extraction, no network, no developer-machine venv path.
    """
    ws = _workspace(tmp_path)
    venv = ws / ".venv"
    subprocess.run(
        [sys.executable, "-m", "venv", "--without-pip", str(venv)],
        check=True,
        capture_output=True,
        text=True,
    )
    venv_py = _venv_python(venv)
    assert venv_py is not None, venv
    purelib = _purelib_of(venv_py)
    (purelib / "pypdf.py").write_text(
        "# capability-selection stub\n", encoding="utf-8",
    )
    (purelib / "openpyxl.py").write_text(
        "# capability-selection stub\n", encoding="utf-8",
    )
    assert not (purelib / "pptx.py").exists()
    assert not (purelib / "pptx").exists()
    assert cebridge.interpreter_has_module(venv_py, "pypdf")
    assert cebridge.interpreter_has_module(venv_py, "openpyxl")
    assert not cebridge.interpreter_has_module(venv_py, "pptx")
    ws_py = str(venv_py)
    assert cebridge.python_for_ingest(ws, ".xlsx") == [ws_py]
    assert cebridge.python_for_ingest(ws, ".pdf") == [ws_py]
    assert cebridge.python_for_ingest(ws, ".pptx") == [sys.executable]
    assert cebridge.script_python_for(ws, "scan.py") == [ws_py]
    assert cebridge.script_python_for(ws, "graph.py") == [ws_py]


def test_real_xlsx_pdf_and_mixed_directory_ingest(tmp_path: Path):
    _require_ce()
    fmt_py = _venv_python(FORMAT_VENV)
    if fmt_py is None:
        pytest.skip("optional format extractor venv unavailable")
    ws = _workspace(tmp_path)
    try:
        (ws / ".venv").symlink_to(FORMAT_VENV, target_is_directory=True)
    except OSError:
        pytest.skip("cannot link format extractor venv")
    raw = ws / "vault" / "raw"
    _write_xlsx(raw / "probe.xlsx", fmt_py)
    _write_text_pdf(raw / "probe.pdf")
    _write_pptx(raw / "probe.pptx")
    (raw / "note.txt").write_text("TXT-MIX-MARKER-9f3c2a17\n", encoding="utf-8")
    xlsx = tools.REGISTRY["ce_ingest"].handler(ws, {"path": "vault/raw/probe.xlsx"})
    pdf = tools.REGISTRY["ce_ingest"].handler(ws, {"path": "vault/raw/probe.pdf"})
    assert ce_tools.ingest_is_success(xlsx), xlsx
    assert ce_tools.ingest_is_success(pdf), pdf
    mixed = tools.REGISTRY["ce_ingest"].handler(ws, {"path": "vault/raw"})
    assert ce_tools.ingest_is_success(mixed), mixed
    body = "\n".join(_extracted_texts(ws))
    assert "XLSX-ROOT-MARKER" in body
    assert "PDF-ROOT-MARKER" in body
    assert PPTX_TITLE in body
    assert "TXT-MIX-MARKER-9f3c2a17" in body
    assert "extraction unavailable" not in body.lower()
    methods = []
    for row in mixed.get("results") or []:
        if isinstance(row, dict) and row.get("extraction_method"):
            methods.append(row["extraction_method"])
    assert "openpyxl" in methods
    assert any(m.startswith("pypdf") for m in methods)
    assert "python-pptx" in methods


def test_real_corrupt_and_good_mixed_directory_ingest(tmp_path: Path):
    _require_ce()
    ws = _workspace(tmp_path)
    raw = ws / "vault" / "raw"
    good = raw / "good.pptx"
    bad = raw / "corrupt.pptx"
    _write_pptx(good)
    bad.write_bytes(b"PK\x03\x04 this is not a real presentation")
    out = tools.REGISTRY["ce_ingest"].handler(ws, {"path": "vault/raw"})
    assert ce_tools.ingest_is_success(out), out
    assert out.get("ok") == 1, out
    assert out.get("failed") == 1, out
    assert good.is_file()
    assert bad.is_file()
    extracts = list((ws / "vault").rglob("*.extracted.md"))
    bodies = [p.read_text(encoding="utf-8", errors="replace") for p in extracts]
    assert any(PPTX_TITLE in b for b in bodies), extracts
    assert not any("corrupt.pptx" in p.name for p in extracts)
    db = ws / "vault" / "vault.db"
    assert db.is_file()
    with sqlite3.connect(db) as conn:
        rows = conn.execute("SELECT path, body FROM sources").fetchall()
    assert any(PPTX_TITLE in (body or "") for _path, body in rows)
    assert not any("corrupt.pptx" in path for path, _body in rows)
    kept_binaries = [
        p for p in (ws / "vault").iterdir()
        if p.is_file() and p.suffix == ".pptx"
    ]
    assert kept_binaries, "original/kept PPTX binaries were removed"


def test_ce_ingest_honors_timeout_payload(tmp_path: Path, monkeypatch: pytest.MonkeyPatch):
    ws = _workspace(tmp_path)
    src = ws / "vault" / "raw" / "a.txt"
    src.write_text("hi", encoding="utf-8")
    seen: dict = {}

    def fake_run(script, args=None, *, cwd, timeout=120.0, require_json=True, python=None):
        seen["timeout"] = timeout
        seen["python"] = python
        return {
            "ok": 1,
            "considered": 1,
            "results": [{
                "ok": True, "extracted": "vault/a.txt.extracted.md",
                "extraction_method": "utf8", "extraction_quality": "good",
            }],
        }

    monkeypatch.setattr(ce_tools.cebridge, "run_script", fake_run)
    (ws / "vault" / "a.txt.extracted.md").write_text("hi\n", encoding="utf-8")
    out = ce_tools._ce_ingest(ws, {"path": "vault/raw/a.txt", "timeout": 90})
    assert 80 <= seen["timeout"] <= 90
    assert out.get("timeout_s") == 90
    assert seen["timeout"] != 300
